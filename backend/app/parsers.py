"""Document parsers. Each returns a list of (text, locator) units."""
from __future__ import annotations

import re
import zipfile

from .errors import UserError

MAX_PDF_PAGES = 500
MAX_TEXT_BYTES = 5 * 1024 * 1024
MAX_UNPACKED_BYTES = 200 * 1024 * 1024  # docx/pptx are zip files: refuse zip bombs


def _clean(t: str) -> str:
    t = t.replace("\x00", "")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def _check_zip(path: str, label: str) -> None:
    try:
        with zipfile.ZipFile(path) as z:
            unpacked = sum(i.file_size for i in z.infolist())
    except (zipfile.BadZipFile, OSError):
        raise UserError(f"Couldn't read this {label}. It may be corrupted.")
    if unpacked > MAX_UNPACKED_BYTES:
        raise UserError(f"This {label} is too large once unpacked.")


def parse_pdf(src) -> list[tuple[str, dict]]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(src)
    except Exception:
        raise UserError("Couldn't read this PDF. It may be corrupted.")
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            raise UserError("This PDF is password protected.")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise UserError(f"PDF has too many pages (limit {MAX_PDF_PAGES}).")
    units = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = _clean(page.extract_text() or "")
        except Exception:
            text = ""
        if text:
            units.append((text, {"kind": "page", "page": i}))
    if not units:
        raise UserError(
            "No readable text found. Scanned PDFs need OCR, which isn't supported yet."
        )
    return units


def parse_docx(path: str) -> list[tuple[str, dict]]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_zip(path, "Word file")
    try:
        doc = Document(path)
    except Exception:
        raise UserError("Couldn't read this Word file. It may be corrupted.")

    units: list[tuple[str, dict]] = []
    section: str | None = None
    buf: list[str] = []

    def flush():
        nonlocal buf
        if buf:
            units.append(("\n".join(buf), {"kind": "section", "section": section}))
        buf = []

    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith("Heading") or style == "Title":
                flush()
                section = text[:120]
            buf.append(text)
        elif tag == "tbl":
            table = Table(el, doc)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    buf.append(" | ".join(cells))
    flush()
    if not units:
        raise UserError("This document has no readable text.")
    return units


def _shape_text(shape) -> list[str]:
    out: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            out.append(t)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            out.extend(_shape_text(s))
    return out


def parse_pptx(path: str) -> list[tuple[str, dict]]:
    from pptx import Presentation

    _check_zip(path, "PowerPoint file")
    try:
        prs = Presentation(path)
    except Exception:
        raise UserError("Couldn't read this PowerPoint file. It may be corrupted.")
    units = []
    for i, slide in enumerate(prs.slides, 1):
        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(_shape_text(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        if parts:
            units.append(("\n".join(parts), {"kind": "slide", "slide": i}))
    if not units:
        raise UserError("This presentation has no readable text.")
    return units


def _decode(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def parse_text(path: str, markdown: bool = False) -> list[tuple[str, dict]]:
    with open(path, "rb") as f:
        raw = f.read(MAX_TEXT_BYTES + 1)
    if len(raw) > MAX_TEXT_BYTES:
        raise UserError("Text file is too large (limit 5 MB).")
    text = _decode(raw)
    if not text.strip():
        raise UserError("This file is empty.")
    if not markdown:
        return [(text, {"kind": "section", "section": None})]

    units: list[tuple[str, dict]] = []
    section, buf = None, []
    for line in text.splitlines():
        m = re.match(r"^#{1,6}\s+(.*\S)\s*$", line)
        if m:
            if buf:
                units.append(("\n".join(buf), {"kind": "section", "section": section}))
            section, buf = m.group(1)[:120], [line]
        else:
            buf.append(line)
    if buf:
        units.append(("\n".join(buf), {"kind": "section", "section": section}))
    return units
