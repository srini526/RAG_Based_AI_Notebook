
import pytest

from app.errors import UserError
from app.parsers import parse_docx, parse_pdf, parse_pptx, parse_text


def test_pdf(tmp_path):
    from fpdf import FPDF
    pdf = FPDF()
    for t in ("Photosynthesis converts light into chemical energy.", "Mitochondria produce ATP."):
        pdf.add_page(); pdf.set_font("Helvetica", size=12); pdf.multi_cell(0, 8, t)
    p = tmp_path / "a.pdf"; pdf.output(str(p))
    units = parse_pdf(str(p))
    assert [u[1]["page"] for u in units] == [1, 2]
    assert "Mitochondria" in units[1][0]


def test_pdf_without_text_is_rejected(tmp_path):
    from fpdf import FPDF
    pdf = FPDF(); pdf.add_page(); p = tmp_path / "blank.pdf"; pdf.output(str(p))
    with pytest.raises(UserError, match="OCR"):
        parse_pdf(str(p))


def test_pdf_garbage(tmp_path):
    p = tmp_path / "x.pdf"; p.write_bytes(b"not a pdf")
    with pytest.raises(UserError):
        parse_pdf(str(p))


def test_docx_sections_and_tables(tmp_path):
    from docx import Document
    d = Document()
    d.add_heading("Intro", 1); d.add_paragraph("Hello world.")
    d.add_heading("Data", 1)
    t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "k"; t.rows[0].cells[1].text = "v"
    p = tmp_path / "a.docx"; d.save(str(p))
    units = parse_docx(str(p))
    assert [u[1]["section"] for u in units] == ["Intro", "Data"]
    assert "k | v" in units[1][0]


def test_pptx_slides_and_notes(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Closures"; s.placeholders[1].text = "Remember scope"
    s.notes_slide.notes_text_frame.text = "mention lexical env"
    p = tmp_path / "a.pptx"; prs.save(str(p))
    units = parse_pptx(str(p))
    assert units[0][1] == {"kind": "slide", "slide": 1}
    assert "Closures" in units[0][0] and "lexical env" in units[0][0]


def test_markdown_headings(tmp_path):
    p = tmp_path / "a.md"; p.write_text("# One\ntext a\n## Two\ntext b\n", encoding="utf-8")
    assert [u[1]["section"] for u in parse_text(str(p), markdown=True)] == ["One", "Two"]


def test_text_latin1_and_empty(tmp_path):
    p = tmp_path / "a.txt"; p.write_bytes("café".encode("latin-1"))
    assert "caf" in parse_text(str(p))[0][0]
    e = tmp_path / "e.txt"; e.write_text("  \n")
    with pytest.raises(UserError):
        parse_text(str(e))
